"""
flow_dashboard.py
-----------------
Streamlit dashboard for adjustment flows.

Upload the adjustment workbook (required) and, optionally, the quantification
workbook. Pick an adjustment in the sidebar and the page shows, top to bottom:

  1. summary cards (steps, inputs, calculations, outputs, file links, issues)
  2. a simple, static flow - Inputs -> Process -> Output - Internal ->
     Output - External as light coloured bands, steps as cards, one arrow
     from each step to the next
  3. the interactive flow (click-for-details, search, arrow checkboxes)
  4. QUANTIFICATION (new, needs the quantification file):
       - metric cards: total records, total data points, number of sheets,
         assigned to
       - a summary panel with assigned to, source systems, totals and sheets
       - a Plotly horizontal bar chart of the seven dimensions
       - a sheet-level table (file, sheet, assigned to, records, grand total)
  5. OVERLAP (new): which steps use a work paper that another adjustment also
     uses, shown on the same flow picture with shared steps outlined, plus a table
  6. CONSOLIDATED FLOW (new): all adjustments on one picture, by stage, with each
     work paper shown once and the adjustments that use it
  7. a table of all steps, and the issues list

Plus sidebar buttons to download, as standalone HTML for all adjustments,
either the Flow overview (light, with Quantification and Overlap tabs) or the
interactive view.

Needs adjustment_flow.py in the SAME folder - it does the reading, arrows,
checks and drawing - and plotly for the quantification chart.

Requires: pip install streamlit openpyxl pandas plotly

Run:
    streamlit run flow_dashboard.py
"""

import difflib
import io
import json
import re
from pathlib import Path

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl import load_workbook

try:
    import plotly.graph_objects as go
except ImportError:          # the page still works; only the quantification chart is skipped
    go = None

st.set_page_config(page_title="Adjustment Flow Dashboard", page_icon="🔀", layout="wide")

# ---------------------------------------------------------------- load adjustment_flow.py
# It must sit in the same folder as this file. Browsers often rename a re-downloaded
# file to "adjustment_flow (1).py", so any adjustment_flow*.py in this folder is accepted.
import importlib.util
import sys

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))
try:
    import adjustment_flow as af
except ModuleNotFoundError:
    candidates = sorted(HERE.glob("adjustment*flow*.py"), key=lambda q: q.stat().st_mtime, reverse=True)
    candidates = [c for c in candidates if c.resolve() != Path(__file__).resolve()]
    if not candidates:
        st.error(f"Can't find **adjustment_flow.py**. Put it in the same folder as this file:\n\n`{HERE}`")
        st.stop()
    spec = importlib.util.spec_from_file_location("adjustment_flow", candidates[0])
    af = importlib.util.module_from_spec(spec)
    sys.modules["adjustment_flow"] = af
    spec.loader.exec_module(af)


st.markdown("""
<style>
  .block-container { padding-top: 1.6rem; padding-bottom: 1rem; }
  div[data-testid="stMetric"] { background: rgba(128,128,128,0.08); border-radius: 8px;
                                padding: 10px 14px; border-top: 4px solid #888; }
  /* quantification metric cards: same look, one accent colour (scoped so the
     flow cards above keep their own colours) */
  .st-key-quant_cards div[data-testid="stMetric"] { border-top-color: #2a9d8f !important; }
</style>
""", unsafe_allow_html=True)

BUCKETS = [("Input", "Inputs", "#2ecc71"), ("Calculation", "Process", "#2a9d8f"),
           ("OutputInternal", "Output - Internal", "#f39c12"),
           ("OutputExternal", "Output - External", "#e76f51")]
OPTIONAL_BUCKETS = [("Output", "Outputs (not split)", "#f4a261"),
                    ("Unassigned", "Unknown stage", "#9ca3af")]
# light styling for the static flow: (band tint, band accent)
BAND_STYLE = {"Input": ("#effaf3", "#2ecc71"), "Calculation": ("#ecf7f6", "#2a9d8f"),
              "OutputInternal": ("#fef6ea", "#f39c12"), "OutputExternal": ("#fdf0ec", "#e76f51"),
              "Output": ("#fef4ec", "#f4a261"), "Unassigned": ("#f3f4f6", "#9ca3af")}
BAND_NAME = {"Input": "Inputs", "Calculation": "Process", "OutputInternal": "Output - Internal",
             "OutputExternal": "Output - External", "Output": "Outputs (not split)",
             "Unassigned": "Unknown stage"}
TYPE_BADGE = {"Temporary": ("#e8f1fb", "#1f6fb2"), "Permanent": ("#f3ebfa", "#7b3fa8")}
ACTION_STYLE = [("download", "#e8f1fb", "#1f6fb2"), ("upload", "#f3ebfa", "#7b3fa8"),
                ("roll", "#fdf5d9", "#8a6d00")]

# ---------------------------------------------------------------- quantification settings
# Header names are matched ignoring case, spaces and punctuation, so "Total Records",
# "total_records" and "TOTAL RECORDS" all count. Change these lists to match the file.
QUANT_FIELDS = {
    "adjustment": ["adjustment", "adjustmentname", "adj"],
    "file": ["filename", "file", "workpapername", "workpaper"],
    "sheet": ["sheetname", "sheet", "worksheet", "tab"],
    "assigned": ["assignedto", "assignee", "owner"],
    "source": ["sourcesystem", "source", "system"],
    "records": ["totalrecords", "records", "recordcount"],
    "grand": ["grandtotal"],
}
# the seven dimensions shown in the chart, in display order
DIMENSIONS = ["Book Reversal", "Book Reclass", "Tax Adjustment", "Book Basis",
              "Tax Basis", "Book Adjustment", "Reconciliation"]
# rows whose Adjustment cell is one of these are totals, not adjustments
TOTAL_ROW_WORDS = {"total", "grandtotal", "subtotal", "sum"}
# "Total Data Points" = the sum of the Grand Total column (per row, then across rows)


# ---------------------------------------------------------------- small helpers
def esc(v):
    return (str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;"))


def clip(text, n):
    text = str(text)
    return text if len(text) <= n else text[:n - 1] + "…"


def wrap2(text, n):
    """Split into at most two lines of ~n chars (second line clipped)."""
    words, lines, cur = str(text).split(), [], ""
    for w in words:
        if len(cur) + len(w) + (1 if cur else 0) <= n:
            cur = f"{cur} {w}".strip()
        else:
            lines.append(cur or w[:n])
            cur = w if cur else ""
    lines.append(cur)
    lines = [l for l in lines if l] or [""]
    return lines if len(lines) <= 2 else [lines[0], clip(" ".join(lines[1:]), n)]


def fmt_amount(v):
    return "-" if v is None else f"{v:,.2f}"


def fmt_num(v):
    """Whole numbers without decimals, others to 2 places."""
    if v is None:
        return "-"
    return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"


def norm(v):
    """Case, spaces and punctuation ignored: 'Total Records' -> 'totalrecords'."""
    return "" if v is None else re.sub(r"[^a-z0-9]", "", str(v).lower())


def to_num(v):
    """Excel cell -> float. Blanks, '-' and text give None (treated as 0 by the caller)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)          # NaN check
    t = str(v).strip().replace(",", "").replace("$", "").replace(" ", "")
    if not t:
        return None
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        n = float(t)
    except ValueError:
        return None
    if n != n:
        return None
    return -n if neg else n


def txt(v):
    """Cell -> clean text. 2024.0 -> '2024'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ---------------------------------------------------------------- static flow (light)
# layout of the static flow (pixels)
CARD_W, CARD_H, GAP_X, ROW_GAP, BAND_GAP = 200, 116, 48, 34, 46
PER_ROW, PAD_X, BAND_TOP, BAND_BOTTOM = 4, 24, 42, 18
FLOW_W = 2 * PAD_X + PER_ROW * CARD_W + (PER_ROW - 1) * GAP_X
FONT = "Source Sans Pro, Segoe UI, Arial, sans-serif"


def static_flow_svg(flow, uid="ah", marks=None):
    """Static flow: light bands top to bottom, one arrow from each step to the next."""
    issue_count = {}
    for i in flow["issues"]:
        if i["id"]:
            issue_count[i["id"]] = issue_count.get(i["id"], 0) + 1
    always = ["Input", "Calculation", "OutputInternal", "OutputExternal"]
    order = [b for b in af.BUCKET_ORDER
             if b in always or any(s["bucket"] == b for s in flow["steps"])]

    bands, pos, y = [], {}, 0
    for b in order:
        in_band = [s for s in flow["steps"] if s["bucket"] == b]
        rows = [in_band[i:i + PER_ROW] for i in range(0, len(in_band), PER_ROW)] or [[]]
        top = y
        cy = top + BAND_TOP
        for row in rows:
            row_w = len(row) * CARD_W + max(0, len(row) - 1) * GAP_X
            x = (FLOW_W - row_w) / 2
            for s in row:
                pos[s["id"]] = (x, cy)
                x += CARD_W + GAP_X
            cy += CARD_H + ROW_GAP
        height = BAND_TOP + len(rows) * CARD_H + (len(rows) - 1) * ROW_GAP + BAND_BOTTOM
        bands.append((b, top, height, len(in_band)))
        y = top + height + BAND_GAP
    total_h = y - BAND_GAP + 4

    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {FLOW_W} {total_h}" '
           f'width="100%" style="max-width:{FLOW_W}px;display:block;font-family:{FONT}">',
           f'<defs><marker id="{uid}" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" '
           'markerHeight="7" orient="auto-start-reverse"><path d="M0,0 L10,5 L0,10 z" '
           'fill="#4b5563"/></marker></defs>']

    for b, top, h, n in bands:
        tint, accent = BAND_STYLE[b]
        out.append(f'<rect x="1" y="{top}" width="{FLOW_W - 2}" height="{h}" rx="12" fill="{tint}"/>')
        out.append(f'<rect x="1" y="{top}" width="6" height="{h}" rx="3" fill="{accent}"/>')
        # title on the right, so arrows coming down into the first card never cross it
        out.append(f'<text x="{FLOW_W - PAD_X}" y="{top + 27}" text-anchor="end" font-size="14" '
                   f'font-weight="700" fill="#374151" letter-spacing="0.6">{BAND_NAME[b].upper()}'
                   f'<tspan font-weight="400" fill="#6b7280" letter-spacing="0">  ·  {n} '
                   f'step{"s" if n != 1 else ""}</tspan></text>')

    # arrows: every step to the next one, drawn under the cards
    steps = flow["steps"]
    for prev, nxt in zip(steps, steps[1:]):
        (x1, y1), (x2, y2) = pos[prev["id"]], pos[nxt["id"]]
        if y1 == y2:   # same row: straight across
            a, b_ = (x1 + CARD_W, x2) if x2 > x1 else (x1, x2 + CARD_W)
            d = f"M{a},{y1 + CARD_H / 2} L{b_ - 2},{y2 + CARD_H / 2}"
        else:          # next row / next band: down, across, down (rounded corners)
            sx, sy, ex, ey = x1 + CARD_W / 2, y1 + CARD_H, x2 + CARD_W / 2, y2 - 2
            my = (sy + ey) / 2
            r = min(10, abs(ex - sx) / 2)
            if r < 1:
                d = f"M{sx},{sy} L{ex},{ey}"
            else:
                dx = 1 if ex > sx else -1
                d = (f"M{sx},{sy} L{sx},{my - r} Q{sx},{my} {sx + dx * r},{my} "
                     f"L{ex - dx * r},{my} Q{ex},{my} {ex},{my + r} L{ex},{ey}")
        out.append(f'<path d="{d}" fill="none" stroke="#4b5563" stroke-width="2" '
                   f'stroke-linejoin="round" marker-end="url(#{uid})"/>')

    for s in steps:
        x, y0 = pos[s["id"]]
        a = (s["action"] or "").lower()
        bg, fg = next(((bb, f) for k, bb, f in ACTION_STYLE if k in a), ("#f3f4f6", "#4b5563"))
        act = clip(s["action"] or "-", 14)
        pw = 8 + 6.2 * len(act)
        tip = "\n".join(f"{k}: {v}" for k, v in s["fields"] if v)
        out.append(f'<g><title>{esc(tip)}</title>')
        out.append(f'<rect x="{x}" y="{y0}" width="{CARD_W}" height="{CARD_H}" rx="10" '
                   f'fill="#ffffff" stroke="#e5e7eb"/>')
        shared = bool(marks) and s["id"] in marks          # overlap view only
        if shared:
            out.append(f'<rect x="{x - 4}" y="{y0 - 4}" width="{CARD_W + 8}" height="{CARD_H + 8}" '
                       'rx="13" fill="none" stroke="#7b3fa8" stroke-width="2" stroke-dasharray="6 4"/>')
        out.append(f'<text x="{x + 12}" y="{y0 + 21}" font-size="11" font-weight="700" '
                   f'fill="#6b7280" letter-spacing="0.5">STEP {esc(s["step"] or "?")}</text>')
        out.append(f'<rect x="{x + CARD_W - 10 - pw}" y="{y0 + 9}" width="{pw}" height="17" '
                   f'rx="8.5" fill="{bg}"/>')
        out.append(f'<text x="{x + CARD_W - 10 - pw / 2}" y="{y0 + 21.5}" font-size="10.5" '
                   f'font-weight="600" fill="{fg}" text-anchor="middle">{esc(act)}</text>')
        for k, line in enumerate(wrap2(s["report"] or "(no report name)", 25)):
            out.append(f'<text x="{x + 12}" y="{y0 + 44 + k * 16}" font-size="13.5" '
                       f'font-weight="600" fill="#111827">{esc(line)}</text>')
        wp = s["location"] or s["source"] or "no work paper"
        if s["location"] and s["wpSheet"]:
            wp += " › " + s["wpSheet"]
        out.append(f'<text x="{x + 12}" y="{y0 + 74}" font-size="11" fill="#6b7280">'
                   f'📄 {esc(clip(wp, 30))}</text>')
        if s["standard"]:
            out.append(f'<text x="{x + 12}" y="{y0 + 90}" font-size="11" fill="#374151">'
                       f'Basis: {esc(clip(s["standard"], 30))}</text>')
        warn = issue_count.get(s["id"])
        if warn:
            out.append(f'<text x="{x + 12}" y="{y0 + 106}" font-size="10.5" font-weight="700" '
                       f'fill="#c0392b">⚠ {warn} issue{"s" if warn > 1 else ""}</text>')
        if shared:
            out.append(f'<text x="{x + CARD_W - 12}" y="{y0 + 106}" font-size="10.5" font-weight="700" '
                       f'fill="#7b3fa8" text-anchor="end">⇄ shared × {len(marks[s["id"]])}</text>')
        out.append('</g>')
    out.append('</svg>')
    return "".join(out), total_h


# ---------------------------------------------------------------- overlap between adjustments
# Two steps overlap when they use the same work paper (file) and the same sheet in it.
# To match on something else, change overlap_key() - e.g. file name only, or report name.
def overlap_key(s):
    loc = norm(s["location"])
    if not loc:
        return ""
    return loc + "|" + norm(s["wpSheet"])


@st.cache_data(show_spinner=False)
def overlap_index(flows):
    """{key: every use of that work paper across all adjustments}."""
    index = {}
    for fi, f in enumerate(flows):
        for s in f["steps"]:
            k = overlap_key(s)
            if not k:
                continue
            wp = (s["location"] or "") + (f" › {s['wpSheet']}" if s["wpSheet"] else "")
            index.setdefault(k, []).append({
                "flow": fi, "name": f["name"], "stage": BAND_NAME.get(s["bucket"], s["bucket"]),
                "step": s["step"], "report": s["report"], "wp": wp})
    return index


def step_overlaps(fi, flow, index):
    """For one adjustment: step id -> the other adjustments' uses of its work paper."""
    marks = {}
    for s in flow["steps"]:
        k = overlap_key(s)
        others = [o for o in index.get(k, []) if o["flow"] != fi] if k else []
        if others:
            marks[s["id"]] = others
    return marks


OVERLAP_CSS = """<style>
  .ov-sum { font-size:15px; margin:0 0 6px; color:#374151; font-family:Source Sans Pro, Segoe UI, Arial, sans-serif; }
  .ov-legend { font-size:13px; color:#6b7280; margin:0 0 10px; font-family:Source Sans Pro, Segoe UI, Arial, sans-serif; }
  .ov-pill { display:inline-block; padding:2px 10px; border-radius:10px; font-size:12.5px; font-weight:600; white-space:nowrap; }
  .ov-yes { background:#f3ebfa; color:#7b3fa8; }
  .ov-no { background:#f3f4f6; color:#6b7280; }
  table.ov { width:100%; border-collapse:collapse; font-size:13.5px; margin-top:14px;
             font-family:Source Sans Pro, Segoe UI, Arial, sans-serif; color:#111827; }
  table.ov th, table.ov td { text-align:left; padding:7px 8px; border-bottom:1px solid #e5e7eb; vertical-align:top; }
  table.ov td:first-child { white-space:nowrap; }
  table.ov th { background:#f9fafb; font-weight:600; color:#374151; }
</style>"""


def overlap_html(fi, flows, index, uid="ov"):
    """Overlap view for one adjustment: the static flow with shared steps outlined, then a table.
    Returns (html, approximate height in pixels for components.html)."""
    flow = flows[fi]
    marks = step_overlaps(fi, flow, index)
    svg, svg_h = static_flow_svg(flow, uid=uid, marks=marks)
    steps = flow["steps"]
    if marks:
        summary = (f'{len(marks)} of {len(steps)} steps use a work paper that another adjustment '
                   f'also uses.')
    else:
        summary = "No step in this adjustment shares a work paper with another adjustment."
    rows = []
    for s in steps:
        others = marks.get(s["id"], [])
        wp = (s["location"] or "-") + (f" › {s['wpSheet']}" if s["location"] and s["wpSheet"] else "")
        if not s["location"]:
            status = '<span class="ov-pill ov-no">No work paper</span>'
        elif others:
            status = '<span class="ov-pill ov-yes">Shared</span>'
        else:
            status = '<span class="ov-pill ov-no">Not shared</span>'
        also = "<br>".join(
            f'{esc(o["name"])} · {esc(o["stage"])} · Step {esc(o["step"] or "?")}'
            + (f' ({esc(o["report"])})' if o["report"] else "")
            for o in others) or "-"
        rows.append(f'<tr><td>Step {esc(s["step"] or "?")}</td><td>{esc(BAND_NAME.get(s["bucket"], s["bucket"]))}</td>'
                    f'<td>{esc(s["report"] or "-")}</td><td>{esc(wp)}</td><td>{status}</td><td>{also}</td></tr>')
    table = ('<table class="ov"><thead><tr><th>Step</th><th>Stage</th><th>Report</th>'
             '<th>Work paper</th><th>Overlap</th><th>Also used in</th></tr></thead>'
             f'<tbody>{"".join(rows)}</tbody></table>')
    html = (OVERLAP_CSS
            + f'<p class="ov-sum">{esc(summary)}</p>'
            + '<p class="ov-legend">Dashed purple outline = the work paper is also used by another '
              'adjustment. Hover a card for its full details.</p>'
            + svg + table)
    return html, int(svg_h) + 150 + 40 * len(steps)


# ---------------------------------------------------------------- overall flow + sheet trace (all adjustments)
ADJ_COLOURS = ["#2a9d8f", "#e76f51", "#7b3fa8", "#1f6fb2", "#c0392b", "#b7791f", "#2e7d32", "#0e7490"]
MAP_MAX_W = 1040                                              # widest the picture gets
MAP_LANE_W, MAP_LANE_GAP, MAP_MARGIN, MAP_LANE_ROW_GAP = 230, 22, 20, 18
MAP_LANE_HEAD, MAP_STEP_H, MAP_STEP_GAP = 50, 76, 14
MAP_STEP_W = MAP_LANE_W - 28
MAP_CARD_W, MAP_CARD_H, MAP_GAP, MAP_ROW_GAP = 200, 60, 22, 16   # input and output cards
MAP_LABEL_H, MAP_BAND_GAP, MAP_PAD = 34, 70, 14
MAP_MAIN = ["Input", "Calculation", "OutputInternal", "OutputExternal"]
MAP_LABEL = {"Input": "Inputs", "Calculation": "Process - one lane per adjustment",
             "OutputInternal": "Output - Internal", "OutputExternal": "Output - External",
             "Output": "Outputs (not split)", "Unassigned": "Unknown stage"}
MAP_CSS = ("<style>"
           ".cons-sum{font-size:15px;margin:0 0 4px;color:#374151;font-family:Source Sans Pro,Segoe UI,Arial,sans-serif}"
           ".cons-legend{font-size:13px;color:#6b7280;margin:0 0 10px;font-family:Source Sans Pro,Segoe UI,Arial,sans-serif}"
           "svg.fm .fm-l,svg.fm .fm-n,svg.fm .fm-a{transition:opacity .15s}"
           "svg.fm .fm-l{opacity:.7}"
           "svg.fm g.fm-a,svg.fm g.fm-n{cursor:pointer}"
           "</style>")
TRACE_STAGE_COLOUR = {"Input": "#2ecc71", "OutputInternal": "#f39c12", "OutputExternal": "#e76f51",
                      "Output": "#f4a261", "Unassigned": "#9ca3af"}


def consolidated_items(flows):
    """Every step of every adjustment, grouped by stage and work paper (report name if none)."""
    groups = {}
    for fi, f in enumerate(flows):
        for s in f["steps"]:
            k = overlap_key(s)
            if not k:
                if not norm(s["report"]):
                    continue
                k = "report|" + norm(s["report"])
            g = groups.setdefault((s["bucket"], k), {
                "stage": s["bucket"], "title": "", "sheet": "", "reports": [], "adjs": {}})
            if s["location"] and not g["title"]:
                g["title"], g["sheet"] = s["location"], s["wpSheet"] or ""
            if s["report"] and s["report"] not in g["reports"]:
                g["reports"].append(s["report"])
            g["adjs"].setdefault(fi, []).append(s["step"])
    for g in groups.values():
        if not g["title"]:
            g["title"] = g["reports"][0] if g["reports"] else "(unnamed)"
    return groups


def overall_flow_html(flows, only=None, uid="fm"):
    """Every adjustment on one picture. Inputs on top (one card per work paper, shared ones once),
    then one process lane per adjustment (lanes wrap into rows of up to four), then internal and
    external outputs. Coloured lines join inputs to the adjustments that use them and adjustments
    to the outputs they produce. Hover or click a lane or card to show only its connections.
    only=<index> draws just that adjustment's flow. Returns (html, height in pixels)."""
    all_groups = consolidated_items(flows)
    lanes = [only] if only is not None else list(range(len(flows)))
    groups = all_groups if only is None else {k: g for k, g in all_groups.items() if only in g["adjs"]}
    n = len(lanes)
    lanes_per_row = max(1, min(n, (MAP_MAX_W - 2 * MAP_MARGIN + MAP_LANE_GAP) // (MAP_LANE_W + MAP_LANE_GAP)))
    lane_row_w = lanes_per_row * MAP_LANE_W + (lanes_per_row - 1) * MAP_LANE_GAP
    width = max(520, min(MAP_MAX_W, 2 * MAP_MARGIN + lane_row_w))
    per_row = max(1, (width - 2 * MAP_MARGIN + MAP_GAP) // (MAP_CARD_W + MAP_GAP))
    stages = MAP_MAIN + [b for b in ("Output", "Unassigned")
                         if any(g["stage"] == b for g in groups.values())]

    # ---- layout
    bands, card_pos, lane_geo, card_idx = [], {}, {}, {}
    y = 0
    for b in stages:
        top = y
        if b == "Calculation":
            row_top = top + MAP_LABEL_H
            for start in range(0, n, lanes_per_row):
                members = lanes[start:start + lanes_per_row]
                row_w = len(members) * MAP_LANE_W + (len(members) - 1) * MAP_LANE_GAP
                x_off = (width - row_w) / 2
                row_h = 0
                for c, fi in enumerate(members):
                    m = max(1, sum(1 for s in flows[fi]["steps"] if s["bucket"] == b))
                    row_h = max(row_h, MAP_LANE_HEAD + 6 + m * MAP_STEP_H + (m - 1) * MAP_STEP_GAP + MAP_PAD)
                    lane_geo[fi] = [x_off + c * (MAP_LANE_W + MAP_LANE_GAP), row_top, 0]
                for fi in members:
                    lane_geo[fi][2] = row_h
                row_top += row_h + MAP_LANE_ROW_GAP
            inner = (row_top - MAP_LANE_ROW_GAP - (top + MAP_LABEL_H)) if n else 24
            count = sum(1 for fi in lanes for s in flows[fi]["steps"] if s["bucket"] == b)
        else:
            items = sorted(((k, g) for k, g in groups.items() if k[0] == b),
                           key=lambda kg: (-len(kg[1]["adjs"]), norm(kg[1]["title"])))
            rows = max(1, -(-len(items) // per_row))
            for idx, (k, g) in enumerate(items):
                r, c = divmod(idx, per_row)
                in_row = min(per_row, len(items) - r * per_row)
                row_w = in_row * MAP_CARD_W + (in_row - 1) * MAP_GAP
                x0 = (width - row_w) / 2 + c * (MAP_CARD_W + MAP_GAP)
                card_pos[k] = (x0, top + MAP_LABEL_H + r * (MAP_CARD_H + MAP_ROW_GAP))
                card_idx[k] = len(card_idx)
            inner = rows * MAP_CARD_H + (rows - 1) * MAP_ROW_GAP if items else 24
            count = len(items)
        bands.append((b, top, MAP_LABEL_H + inner + MAP_PAD, count))
        y = top + MAP_LABEL_H + inner + MAP_PAD + MAP_BAND_GAP
    total_h = y - MAP_BAND_GAP + 6

    def col(fi):
        return ADJ_COLOURS[fi % len(ADJ_COLOURS)]

    # ---- hover and click-to-lock focus rules (scoped to this picture's svg)
    focus = []
    for fi in lanes:
        for state in (f'svg.fm:has(.fm-a{fi}:hover)', f'svg.fm[data-focus="a{fi}"]'):
            focus.append(f"{state} .fm-l:not(.fm-l{fi}),{state} .fm-n:not(.fm-c{fi}),"
                         f"{state} .fm-a:not(.fm-a{fi}){{opacity:.12}}")
    for k, i in card_idx.items():
        for state in (f'svg.fm:has(.fm-k{i}:hover)', f'svg.fm[data-focus="k{i}"]'):
            focus.append(f"{state} .fm-l:not(.fm-k{i}),{state} .fm-n:not(.fm-k{i}){{opacity:.12}}")

    out = [f'<svg class="fm" id="{uid}" xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {total_h}" '
           f'width="100%" style="max-width:{width}px;display:block;font-family:{FONT}"><defs>']
    for fi in lanes:
        out.append(f'<marker id="{uid}c{fi}" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" '
                   f'markerHeight="7" orient="auto-start-reverse"><path d="M0,0 L10,5 L0,10 z" '
                   f'fill="{col(fi)}"/></marker>')
    out.append(f'<marker id="{uid}k" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="6" '
               f'markerHeight="6" orient="auto-start-reverse"><path d="M0,0 L10,5 L0,10 z" '
               f'fill="#4b5563"/></marker></defs>')
    out.append("<style>" + "".join(focus) + "</style>")

    # ---- bands
    for b, top, h, count in bands:
        tint, accent = BAND_STYLE.get(b, ("#f3f4f6", "#9ca3af"))
        out.append(f'<rect x="2" y="{top}" width="{width - 4}" height="{h}" rx="12" fill="{tint}"/>')
        out.append(f'<rect x="2" y="{top}" width="6" height="{h}" rx="3" fill="{accent}"/>')
        sub = "" if b == "Calculation" else f'  ·  {count} item{"s" if count != 1 else ""}'
        out.append(f'<text x="{MAP_MARGIN}" y="{top + 23}" font-size="13" font-weight="700" '
                   f'fill="#374151" letter-spacing="0.6">{esc(MAP_LABEL.get(b, b).upper())}'
                   f'<tspan font-weight="400" fill="#6b7280" letter-spacing="0">{esc(sub)}</tspan></text>')
        if b != "Calculation" and count == 0:
            out.append(f'<text x="{width / 2}" y="{top + MAP_LABEL_H + 16}" font-size="12" '
                       f'fill="#6b7280" text-anchor="middle">none in this stage</text>')

    # ---- connecting lines (under the cards and lanes)
    for (stage, key), g in groups.items():
        if (stage, key) not in card_pos:
            continue
        x, y0 = card_pos[(stage, key)]
        ci = card_idx[(stage, key)]
        for fi in sorted(g["adjs"]):
            if fi not in lane_geo:
                continue
            lx, lt, lh = lane_geo[fi]
            lcx = lx + MAP_LANE_W / 2
            if stage == "Input":
                x1, y1, x2, y2 = x + MAP_CARD_W / 2, y0 + MAP_CARD_H, lcx, lt
            else:
                x1, y1, x2, y2 = lcx, lt + lh, x + MAP_CARD_W / 2, y0
            my = (y1 + y2) / 2
            out.append(f'<path class="fm-l fm-l{fi} fm-k{ci}" d="M{x1},{y1} C{x1},{my} {x2},{my} {x2},{y2}" '
                       f'fill="none" stroke="{col(fi)}" stroke-width="1.8" '
                       f'marker-end="url(#{uid}c{fi})"/>')

    # ---- process lanes, one per adjustment
    for fi in lanes:
        f = flows[fi]
        lx, lt, lh = lane_geo[fi]
        steps = [s for s in f["steps"] if s["bucket"] == "Calculation"]
        out.append(f'<g class="fm-a fm-a{fi}" data-fi="{fi}">')
        out.append(f'<rect x="{lx}" y="{lt}" width="{MAP_LANE_W}" height="{lh}" rx="12" '
                   f'fill="#ffffff" stroke="{col(fi)}" stroke-width="1.5"/>')
        out.append(f'<rect x="{lx}" y="{lt}" width="{MAP_LANE_W}" height="{MAP_LANE_HEAD}" rx="12" '
                   f'fill="{col(fi)}"/>')
        out.append(f'<text x="{lx + 12}" y="{lt + 21}" font-size="13.5" font-weight="700" fill="#fff">'
                   f'{esc(clip(f["name"], 26))}</text>')
        out.append(f'<text x="{lx + 12}" y="{lt + 38}" font-size="11" fill="#fff" opacity="0.9">'
                   f'{esc((f["adjType"] or "Adjustment") + " adjustment")}</text>')
        if not steps:
            out.append(f'<text x="{lx + MAP_LANE_W / 2}" y="{lt + MAP_LANE_HEAD + 30}" font-size="12" '
                       f'fill="#6b7280" text-anchor="middle">no process steps</text>')
        for k, s in enumerate(steps):
            sx = lx + 14
            sy = lt + MAP_LANE_HEAD + 6 + k * (MAP_STEP_H + MAP_STEP_GAP)
            wp = (s["location"] or "no work paper") + (
                f" › {s['wpSheet']}" if s["location"] and s["wpSheet"] else "")
            out.append(f'<rect x="{sx}" y="{sy}" width="{MAP_STEP_W}" height="{MAP_STEP_H}" rx="8" '
                       f'fill="#ffffff" stroke="#e5e7eb"/>')
            out.append(f'<text x="{sx + 10}" y="{sy + 17}" font-size="10.5" font-weight="700" '
                       f'fill="#6b7280" letter-spacing="0.4">STEP {esc(s["step"] or "?")}'
                       f' · {esc(clip((s["action"] or "").upper(), 12))}</text>')
            out.append(f'<text x="{sx + 10}" y="{sy + 35}" font-size="13" font-weight="600" fill="#111827">'
                       f'{esc(clip(s["report"] or "(no name)", 26))}</text>')
            out.append(f'<text x="{sx + 10}" y="{sy + 51}" font-size="10.5" fill="#6b7280">'
                       f'📄 {esc(clip(wp, 28))}</text>')
            if s["standard"]:
                out.append(f'<text x="{sx + 10}" y="{sy + 66}" font-size="10.5" fill="#374151">'
                           f'Basis: {esc(clip(s["standard"], 26))}</text>')
            if k < len(steps) - 1:
                cx = lx + MAP_LANE_W / 2
                out.append(f'<line x1="{cx}" y1="{sy + MAP_STEP_H}" x2="{cx}" '
                           f'y2="{sy + MAP_STEP_H + MAP_STEP_GAP}" stroke="#4b5563" '
                           f'stroke-width="1.6" marker-end="url(#{uid}k)"/>')
        out.append("</g>")

    # ---- input and output cards
    for (stage, key), g in groups.items():
        if (stage, key) not in card_pos:
            continue
        x, y0 = card_pos[(stage, key)]
        ci = card_idx[(stage, key)]
        shared = len(g["adjs"]) > 1
        edge = "#7b3fa8" if shared else "#e5e7eb"
        dash = ' stroke-dasharray="6 4"' if shared else ""
        sub = f"› {g['sheet']}" if g["sheet"] else "; ".join(g["reports"])
        classes = " ".join(["fm-n", f"fm-k{ci}"] + [f"fm-c{fi}" for fi in sorted(g["adjs"])])
        out.append(f'<g class="{classes}" data-k="{ci}">')
        out.append(f'<rect x="{x}" y="{y0}" width="{MAP_CARD_W}" height="{MAP_CARD_H}" rx="10" '
                   f'fill="#ffffff" stroke="{edge}" stroke-width="{2 if shared else 1}"{dash}/>')
        out.append(f'<text x="{x + 12}" y="{y0 + 22}" font-size="13.5" font-weight="600" fill="#111827">'
                   f'{esc(clip(g["title"], 28))}</text>')
        out.append(f'<text x="{x + 12}" y="{y0 + 39}" font-size="11" fill="#6b7280">'
                   f'{esc(clip(sub, 30))}</text>')
        label = f"⇄ shared by {len(g['adjs'])}" if shared else "used by 1"
        out.append(f'<text x="{x + 12}" y="{y0 + MAP_CARD_H - 10}" font-size="10.5" font-weight="700" '
                   f'fill="{"#7b3fa8" if shared else "#6b7280"}">{esc(label)}</text>')
        for j, fi in enumerate(sorted(g["adjs"])):
            cx = x + MAP_CARD_W - 12 - j * 14
            out.append(f'<circle cx="{cx}" cy="{y0 + MAP_CARD_H - 14}" r="5" fill="{col(fi)}">'
                       f'<title>{esc(flows[fi]["name"])}</title></circle>')
        out.append("</g>")
    out.append("</svg>")

    def cnt(stage):
        items = [g for (s_, _), g in groups.items() if s_ == stage]
        return f"{len(items)} ({sum(1 for g in items if len(g['adjs']) > 1)} shared)"

    if only is None:
        summary = (f"{n} adjustment{'s' if n != 1 else ''}. Inputs: {cnt('Input')}. "
                   f"Output - Internal: {cnt('OutputInternal')}. Output - External: {cnt('OutputExternal')}.")
    else:
        summary = (f"{flows[only]['name']}: {cnt('Input')} inputs, {cnt('OutputInternal')} internal outputs, "
                   f"{cnt('OutputExternal')} external outputs (counts in brackets are shared with other adjustments).")
    legend = ("Click a lane or card to lock its connections (click again or the background to release). "
              "Hover to preview. Dashed purple = used by more than one adjustment.")
    html = (MAP_CSS + f'<p class="cons-sum">{esc(summary)}</p>'
            + f'<p class="cons-legend">{esc(legend)}</p>' + "".join(out))
    return html, int(total_h) + 90


# click-to-lock: one delegated listener for every overall-flow picture on the page
FM_JS = r"""
(function () {
  document.addEventListener("click", function (e) {
    var t = e.target;
    var svg = t && t.closest ? t.closest("svg.fm") : null;
    if (!svg) return;
    var g = t.closest("g.fm-a, g.fm-n");
    if (!g) { svg.removeAttribute("data-focus"); return; }
    var key = g.classList.contains("fm-a") ? "a" + g.getAttribute("data-fi") : "k" + g.getAttribute("data-k");
    if (svg.getAttribute("data-focus") === key) svg.removeAttribute("data-focus");
    else svg.setAttribute("data-focus", key);
  });
})();
"""


def fm_sections(flows):
    """Overall flow tab: a dropdown (all adjustments, or one at a time) and one picture per choice."""
    opts = ['<option value="all">All adjustments</option>'] + [
        f'<option value="{i}">{esc(f["name"])}</option>' for i, f in enumerate(flows)]
    secs = [f'<section class="fmsec on" data-key="all">{overall_flow_html(flows, None, uid="fmall")[0]}</section>']
    for i in range(len(flows)):
        secs.append(f'<section class="fmsec" data-key="{i}">'
                    f'{overall_flow_html(flows, i, uid=f"fm{i}")[0]}</section>')
    return ('<div class="pick-row"><label for="fmPick"><b>Show</b></label>'
            f'<select id="fmPick">{"".join(opts)}</select></div>'
            '<p class="fm-caption">Pick one adjustment to see only its flow. Work papers it shares with '
            'other adjustments show the other adjustments as coloured dots.</p>' + "".join(secs))


# ---- sheet trace view: the same dataflow in the trace tool's format (vis-network, dark)
TRACE_CSS = (".trc-root{background:#111;color:#fff;font-family:sans-serif;font-size:13px;padding:10px;"
             "border-radius:8px}"
             ".trc-bar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin-bottom:8px}"
             ".trc-search,.trc-num{background:#222;color:#fff;border:1px solid #555;border-radius:4px;"
             "padding:4px 6px;font-size:13px}"
             ".trc-num{width:60px}"
             ".trc-btn{background:#333;color:#fff;border:1px solid #555;border-radius:4px;padding:4px 10px;"
             "cursor:pointer;font-size:13px}"
             ".trc-status{opacity:.8}"
             ".trc-panel{background:rgba(0,0,0,.6);border:1px solid #333;border-radius:6px;padding:8px;"
             "margin-bottom:8px;max-height:200px;overflow:auto}"
             ".trc-panelbtns{margin-bottom:6px}"
             ".trc-list label{display:inline-block;min-width:200px;margin:2px 8px 2px 0;white-space:nowrap}"
             ".trc-canvas{height:680px;border:1px solid #222;border-radius:6px}"
             ".trc-hint{opacity:.7;font-size:12px;margin-top:6px}")

TRACE_JS = r"""
var TRC_N = 0;
var TRACE_OPTS = {
  layout: {hierarchical: {enabled: true, direction: "UD", sortMethod: "directed",
                          levelSeparation: 150, nodeSpacing: 150, treeSpacing: 120}},
  physics: {hierarchicalRepulsion: {nodeDistance: 170}, solver: "hierarchicalRepulsion"},
  edges: {smooth: {type: "cubicBezier", forceDirection: "vertical"}, arrows: {to: {enabled: true}},
          color: {color: "#888888"}, font: {color: "#ffffff", size: 12, strokeWidth: 0}},
  nodes: {font: {color: "#ffffff", size: 14, multi: true}, borderWidth: 1}
};
var TRACE_HTML = '<div class="trc-bar"><input class="trc-search" list="LISTID" placeholder="Search...">'
  + '<datalist id="LISTID"></datalist>'
  + '<button class="trc-btn" data-a="find">Find</button>'
  + '<label>Min in &ge; <input class="trc-num" type="number" min="0" value="0" data-a="minIn"></label>'
  + '<label>Min out &ge; <input class="trc-num" type="number" min="0" value="0" data-a="minOut"></label>'
  + '<button class="trc-btn" data-a="reset">Reset</button>'
  + '<button class="trc-btn" data-a="list">Sheets</button>'
  + '<span class="trc-status"></span></div>'
  + '<div class="trc-panel" style="display:none"><div class="trc-panelbtns">'
  + '<button class="trc-btn" data-a="all">Show all</button> <button class="trc-btn" data-a="none">Hide all</button> '
  + '<button class="trc-btn" data-a="undo">Undo</button></div><div class="trc-list"></div></div>'
  + '<div class="trc-canvas"></div>'
  + '<div class="trc-hint">Click a box to highlight what it connects to (click again or the background to clear). '
  + 'Right-click to hide or show it. Green = input, orange and red = outputs, coloured = adjustments.</div>';

function mountTrace(root, nodeList, edgeList) {
  var uid = "trc" + (++TRC_N);
  root.innerHTML = TRACE_HTML.split("LISTID").join(uid + "dl");
  var q = function (s) { return root.querySelector(s); };
  var base = {};
  var nodes = new vis.DataSet(nodeList.map(function (n) {
    base[n.id] = n.color; return Object.assign({}, n); }));
  var edges = new vis.DataSet(edgeList.map(function (e) {
    return Object.assign({}, e, {color: {color: e.ec}}); }));
  var network = new vis.Network(q(".trc-canvas"), {nodes: nodes, edges: edges}, TRACE_OPTS);
  network.once("stabilizationIterationsDone", function () {
    network.setOptions({physics: false});
    network.fit({animation: false});
  });
  setTimeout(function () { network.redraw(); network.fit({animation: false}); }, 350);
  var ids = nodes.getIds(), checks = {}, undoStack = [], focusId = null;
  root.__net = network;

  function applyChanges(changes) {
    var entry = [];
    changes.forEach(function (c) {
      var prev = !!nodes.get(c.id).hidden;
      if (prev === !!c.hidden) return;
      entry.push({id: c.id, prev: prev});
      nodes.update({id: c.id, hidden: c.hidden});
      if (checks[c.id]) checks[c.id].checked = !c.hidden;
    });
    if (entry.length) undoStack.push(entry);
  }

  function clearFocus() {
    focusId = null;
    nodes.update(ids.map(function (id) { return {id: id, color: base[id]}; }));
    edges.update(edges.getIds().map(function (eid) {
      return {id: eid, color: {color: edges.get(eid).ec}}; }));
  }

  function setFocus(id) {
    if (focusId === id) { clearFocus(); return; }
    clearFocus();
    focusId = id;
    var keep = {};
    keep[id] = true;
    network.getConnectedNodes(id).forEach(function (x) { keep[x] = true; });
    nodes.update(ids.map(function (x) {
      return keep[x] ? {id: x, color: base[x]} : {id: x, color: "#2a2a2a"}; }));
    edges.update(edges.getIds().map(function (eid) {
      var e = edges.get(eid);
      var on = (e.from === id || e.to === id);
      return {id: eid, color: {color: on ? e.ec : "#2e2e2e"}};
    }));
  }

  network.on("click", function (p) {
    if (p.nodes.length) setFocus(p.nodes[0]); else clearFocus();
  });
  network.on("oncontext", function (p) {
    p.event.preventDefault();
    var id = network.getNodeAt(p.pointer.DOM);
    if (id) applyChanges([{id: id, hidden: !nodes.get(id).hidden}]);
  });

  var list = q(".trc-list");
  ids.slice().sort(function (a, b) {
    return nodes.get(a).searchText.localeCompare(nodes.get(b).searchText); }).forEach(function (id) {
    var lab = document.createElement("label"), cb = document.createElement("input");
    cb.type = "checkbox";
    cb.checked = true;
    cb.addEventListener("change", function () { applyChanges([{id: id, hidden: !cb.checked}]); });
    lab.appendChild(cb);
    lab.appendChild(document.createTextNode(" " + nodes.get(id).searchText));
    list.appendChild(lab);
    checks[id] = cb;
    var opt = document.createElement("option");
    opt.value = nodes.get(id).searchText;
    q("datalist").appendChild(opt);
  });
  q("datalist").id = uid + "dl";

  function applyThresholds() {
    var mi = parseInt(q('[data-a="minIn"]').value, 10) || 0;
    var mo = parseInt(q('[data-a="minOut"]').value, 10) || 0;
    if (mi === 0 && mo === 0) {
      applyChanges(ids.map(function (id) { return {id: id, hidden: false}; }));
      return;
    }
    applyChanges(ids.map(function (id) {
      var n = nodes.get(id);
      var pass = (mi > 0 && n.incoming >= mi) || (mo > 0 && n.outgoing >= mo);
      return {id: id, hidden: !pass};
    }));
  }

  function runSearch() {
    var v = (q(".trc-search").value || "").trim().toLowerCase();
    var st = q(".trc-status");
    if (!v) return;
    var m = ids.find(function (id) { return nodes.get(id).searchText.toLowerCase() === v; })
         || ids.find(function (id) { return nodes.get(id).searchText.toLowerCase().indexOf(v) !== -1; });
    if (!m) { st.textContent = 'No match for "' + v + '"'; return; }
    applyChanges([{id: m, hidden: false}]);
    setFocus(m);
    network.focus(m, {scale: 1.2, animation: {duration: 400, easingFunction: "easeInOutQuad"}});
    st.textContent = "Found: " + nodes.get(m).searchText;
  }

  root.addEventListener("click", function (ev) {
    var a = ev.target.getAttribute ? ev.target.getAttribute("data-a") : null;
    if (a === "list") { var p = q(".trc-panel"); p.style.display = p.style.display === "none" ? "block" : "none"; }
    if (a === "all") applyChanges(ids.map(function (id) { return {id: id, hidden: false}; }));
    if (a === "none") applyChanges(ids.map(function (id) { return {id: id, hidden: true}; }));
    if (a === "undo") {
      var e = undoStack.pop();
      if (e) e.forEach(function (c) {
        nodes.update({id: c.id, hidden: c.prev});
        if (checks[c.id]) checks[c.id].checked = !c.prev;
      });
    }
    if (a === "reset") { q('[data-a="minIn"]').value = 0; q('[data-a="minOut"]').value = 0; applyThresholds(); }
    if (a === "find") runSearch();
  });
  root.addEventListener("input", function (ev) {
    var a = ev.target.getAttribute ? ev.target.getAttribute("data-a") : null;
    if (a === "minIn" || a === "minOut") applyThresholds();
  });
  q(".trc-search").addEventListener("keydown", function (e) { if (e.key === "Enter") runSearch(); });
}
"""


def trace_graph(flows, only=None):
    """Nodes and edges for the sheet-trace view: inputs -> adjustments -> outputs.
    Every edge is labelled with the number of steps it stands for."""
    groups = consolidated_items(flows)
    lanes = [only] if only is not None else list(range(len(flows)))
    lane_set = set(lanes)
    nodes, edges = [], []
    arrows_in = {fi: 0 for fi in lanes}
    arrows_out = {fi: 0 for fi in lanes}
    cards = []
    for idx, ((stage, key), g) in enumerate(sorted(groups.items(), key=lambda kg: (kg[0][0], -len(kg[1]["adjs"]), norm(kg[1]["title"])))):
        use = [fi for fi in sorted(g["adjs"]) if fi in lane_set]
        if not use:
            continue
        is_in = stage == "Input"
        nid = f"c{idx}"
        title = g["title"] + (f" › {g['sheet']}" if g["sheet"] else "")
        detail = "\n".join(f'{flows[fi]["name"]}: step {", ".join(str(x or "?") for x in g["adjs"][fi])}'
                           for fi in use)
        nodes.append({
            "id": nid, "label": f"{clip(title, 34)}\n{len(use)} adjustment{'s' if len(use) != 1 else ''}",
            "color": TRACE_STAGE_COLOUR.get(stage, "#9ca3af"), "shape": "box",
            "searchText": title, "title": f"{title}\n{detail}",
            "incoming": 0 if is_in else len(use), "outgoing": len(use) if is_in else 0})
        cards.append((nid, is_in, use, g))
        for fi in use:
            if is_in:
                arrows_in[fi] += 1
            else:
                arrows_out[fi] += 1
    max_count = 1
    raw = []
    for nid, is_in, use, g in cards:
        for fi in use:
            cnt = len(g["adjs"][fi])
            max_count = max(max_count, cnt)
            raw.append((nid, fi, is_in, cnt))
    for nid, fi, is_in, cnt in raw:
        src, dst = (nid, f"a{fi}") if is_in else (f"a{fi}", nid)
        col = TRACE_STAGE_COLOUR["Input"] if is_in else "#8a8a8a"
        edges.append({"from": src, "to": dst, "width": 1 + cnt / max_count * 6, "label": str(cnt),
                      "title": f"{cnt} step{'s' if cnt != 1 else ''}", "ec": "#8a8a8a"})
    for fi in lanes:
        f = flows[fi]
        n_proc = sum(1 for s in f["steps"] if s["bucket"] == "Calculation")
        nodes.append({
            "id": f"a{fi}",
            "label": f"{clip(f['name'], 30)}\n{f['adjType'] or 'Adjustment'} · {n_proc} process step{'s' if n_proc != 1 else ''}",
            "color": ADJ_COLOURS[fi % len(ADJ_COLOURS)], "shape": "box", "searchText": f["name"],
            "title": f"{f['name']}\nInputs: {arrows_in[fi]}  ·  Outputs: {arrows_out[fi]}",
            "incoming": arrows_in[fi], "outgoing": arrows_out[fi]})
    return nodes, edges


def _json_for_script(obj):
    return json.dumps(obj).replace("</", "<\\/")


@st.cache_data(show_spinner=False)
def trace_doc(flows, only=None):
    """Standalone sheet-trace page for the Streamlit view (vis-network embedded, works offline)."""
    nodes, edges = trace_graph(flows, only)
    data = _json_for_script({"nodes": nodes, "edges": edges})
    return ("<!doctype html><html><head><meta charset='utf-8'>"
            "<style>html,body{margin:0;background:#111}" + TRACE_CSS + "</style>"
            "<script>" + af.get_vis_js() + "</script></head><body>"
            '<div class="trc-root" id="root"></div>'
            "<script>" + TRACE_JS + "\nvar D = " + data + ";\n"
            "mountTrace(document.getElementById('root'), D.nodes, D.edges);</script>"
            "</body></html>")


def trace_sections(flows):
    """Sheet trace tab for the downloadable overview: dropdown + one lazily-drawn map per choice."""
    keys = ["all"] + [str(i) for i in range(len(flows))]
    data, secs = {}, []
    for key in keys:
        only = None if key == "all" else int(key)
        nodes, edges = trace_graph(flows, only)
        data[key] = {"nodes": nodes, "edges": edges}
        secs.append(f'<section class="trsec{" on" if key == "all" else ""}" data-key="{key}">'
                    f'<div class="trc-root" data-key="{key}"></div></section>')
    opts = ['<option value="all">All adjustments</option>'] + [
        f'<option value="{i}">{esc(f["name"])}</option>' for i, f in enumerate(flows)]
    html = (f'<style>{TRACE_CSS}</style>'
            '<div class="pick-row"><label for="trPick"><b>Show</b></label>'
            f'<select id="trPick">{"".join(opts)}</select></div>' + "".join(secs))
    return html, _json_for_script(data)


PAGE_JS = (r"""
var TRACE_DATA = __TRDATA__;
""" + TRACE_JS + FM_JS + r"""
function showSection(cls, key) {
  document.querySelectorAll(cls).forEach(function (s) {
    s.classList.toggle("on", s.getAttribute("data-key") === key);
  });
}
function mountTraceFor(key) {
  var root = document.querySelector('.trc-root[data-key="' + key + '"]');
  if (!root || root.getAttribute("data-mounted")) return;
  root.setAttribute("data-mounted", "1");
  mountTrace(root, TRACE_DATA[key].nodes, TRACE_DATA[key].edges);
}
document.getElementById("fmPick").addEventListener("change", function () {
  showSection(".fmsec", this.value);
});
document.getElementById("trPick").addEventListener("change", function () {
  showSection(".trsec", this.value);
  mountTraceFor(this.value);
});
function onViewChange(v) {
  if (v === "trace") mountTraceFor(document.getElementById("trPick").value);
}
""")



# ---------------------------------------------------------------- reading the adjustment workbook
class NamedBytes(io.BytesIO):
    """In-memory Excel file that keeps its file name (the loader uses .name)."""
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name


@st.cache_data(show_spinner="Reading tables...")
def parse(files, sheet):
    """files: tuple of (name, bytes). Cached, so switching flows is instant."""
    handles = [NamedBytes(data, name) for name, data in files]
    return af.build_payload(af.load_flows(handles, sheet or None))


# ---------------------------------------------------------------- reading the quantification workbook
def quant_header(row):
    """If this row is a quantification header, return (column map, dimension map)."""
    colmap, dimmap = {}, {}
    for idx, v in enumerate(row):
        key = norm(v)
        if not key:
            continue
        for field, aliases in QUANT_FIELDS.items():
            if field not in colmap and key in aliases:
                colmap[field] = idx
                break
        for d in DIMENSIONS:
            if d not in dimmap and key == norm(d):
                dimmap[d] = idx
                break
    # a real header has Adjustment plus at least three other recognised columns
    if "adjustment" in colmap and len(colmap) + len(dimmap) >= 4:
        return colmap, dimmap
    return None


def cell(row, i):
    return row[i] if i is not None and i < len(row) else None


@st.cache_data(show_spinner="Reading quantification file...")
def parse_quant(data):
    """Every quantification row on every sheet. Several rows per adjustment are normal.

    Returns {"rows": [...], "dims_present": set, "found_header": bool}. Missing or
    non-numeric values become 0 (or "" for text), so one blank cell never breaks the page.
    """
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows, dims_present, found_header = [], set(), False
    try:
        for ws in wb.worksheets:
            colmap, dimmap = None, {}
            for r_no, raw in enumerate(ws.iter_rows(values_only=True), start=1):
                hdr = quant_header(raw)
                if hdr:
                    colmap, dimmap = hdr
                    found_header = True
                    dims_present.update(dimmap)
                    continue
                if colmap is None:
                    continue
                adj = txt(cell(raw, colmap.get("adjustment")))
                if not adj or norm(adj) in TOTAL_ROW_WORDS:
                    continue
                rows.append({
                    "adjustment": adj,
                    "adjKey": norm(adj),
                    "file": txt(cell(raw, colmap.get("file"))),
                    "sheet": txt(cell(raw, colmap.get("sheet"))),
                    "assigned": txt(cell(raw, colmap.get("assigned"))),
                    "source": txt(cell(raw, colmap.get("source"))),
                    "records": to_num(cell(raw, colmap.get("records"))) or 0.0,
                    "grand": to_num(cell(raw, colmap.get("grand"))) or 0.0,
                    "dims": {d: (to_num(cell(raw, dimmap[d])) or 0.0) if d in dimmap else None
                             for d in DIMENSIONS},
                    "sheet_in_book": ws.title,
                    "row": r_no,
                })
    finally:
        wb.close()
    return {"rows": rows, "dims_present": dims_present, "found_header": found_header}


def candidate_names(flow):
    """Names the selected flow might go by in the quantification file."""
    name = flow["name"]
    names = [name]
    if " - " in name:                                  # "file - adjustment" (several workbooks)
        names.append(name.split(" - ", 1)[1])
    # "Sheet (2)" when an adjustment had no name of its own
    names = names + [re.sub(r"\s*\(\d+\)$", "", n) for n in names]
    return list(dict.fromkeys(n for n in names if n))


def match_adjustment(flow, quant):
    """Find this flow's rows in the quantification file. Exact match first, then similar name."""
    keys = {}
    for r in quant["rows"]:
        keys.setdefault(r["adjKey"], r["adjustment"])
    for n in candidate_names(flow):
        if norm(n) in keys:
            return keys[norm(n)], "exact"
    for n in candidate_names(flow):
        close = difflib.get_close_matches(norm(n), list(keys), n=1, cutoff=0.8)
        if close:
            return keys[close[0]], "similar"
    return None, None


def summarise(rows, dims_present):
    """Aggregate every row of one adjustment."""
    sheets = {(r["file"], r["sheet"]) for r in rows if r["file"] or r["sheet"]}
    assigned = sorted({r["assigned"] for r in rows if r["assigned"]})
    sources = sorted({r["source"] for r in rows if r["source"]})
    dims = {d: sum((r["dims"].get(d) or 0.0) for r in rows) for d in DIMENSIONS}
    return {
        "records": sum(r["records"] for r in rows),
        "datapoints": sum(r["grand"] for r in rows),
        "sheets": len(sheets),
        "assigned": assigned,
        "sources": sources,
        "dims": dims,
        "missing_dims": [d for d in DIMENSIONS if d not in dims_present],
        "zero_dims": [d for d in DIMENSIONS if d in dims_present and not dims[d]],
    }


def sheet_table(rows):
    """One line per file + sheet, with the numbers added up."""
    grouped = {}
    for r in rows:
        g = grouped.setdefault((r["file"], r["sheet"]), {
            "File Name": r["file"], "Sheet Name": r["sheet"], "assigned": [],
            "Total Records": 0.0, "Grand Total": 0.0})
        g["Total Records"] += r["records"]
        g["Grand Total"] += r["grand"]
        if r["assigned"] and r["assigned"] not in g["assigned"]:
            g["assigned"].append(r["assigned"])
    return pd.DataFrame([{
        "File Name": g["File Name"], "Sheet Name": g["Sheet Name"],
        "Assigned To": ", ".join(g["assigned"]),
        "Total Records": g["Total Records"], "Grand Total": g["Grand Total"],
    } for g in grouped.values()])


def dimension_figure(dims, dims_present):
    """Horizontal bars, one per dimension, in the order of DIMENSIONS."""
    labels, values, colours, texts = [], [], [], []
    for d in DIMENSIONS:
        v = dims.get(d) or 0.0
        present = d in dims_present
        labels.append(d if present else f"{d} (not in file)")
        values.append(v)
        colours.append("#2a9d8f" if v else "#cbd5e1")          # grey = zero or missing
        texts.append(fmt_num(v) if v else "0")
    fig = go.Figure(go.Bar(
        y=labels, x=values, orientation="h", text=texts, textposition="outside",
        cliponaxis=False, marker_color=colours,
        hovertemplate="%{y}<br>%{x:,.2f}<extra></extra>"))
    fig.update_yaxes(autorange="reversed", automargin=True, title=None)
    fig.update_xaxes(title=None, gridcolor="#eef2f7", zeroline=True, rangemode="tozero")
    fig.update_layout(template="plotly_white", showlegend=False, height=70 + 46 * len(DIMENSIONS),
                      margin=dict(l=10, r=70, t=10, b=10))
    return fig


# ---------------------------------------------------------------- standalone HTML pages
OVERVIEW_PAGE = """<!doctype html>
<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Adjustment Flow Overview</title>
<script>__VIS__</script>
<style>
  body { margin:0; background:#fff; color:#111827; font-family:Source Sans Pro, Segoe UI, Arial, sans-serif; }
  header { position:sticky; top:0; z-index:5; background:#f9fafb; border-bottom:1px solid #e5e7eb;
           padding:12px 24px; display:flex; align-items:center; gap:12px; flex-wrap:wrap; }
  header h1 { font-size:18px; margin:0 12px 0 0; }
  select { font-size:15px; padding:7px 10px; min-width:320px; border:1px solid #d1d5db; border-radius:6px;
           background:#fff; cursor:pointer; }
  .nav { cursor:pointer; font-size:22px; color:#6b7280; padding:0 6px; user-select:none; }
  .nav:hover { color:#111827; }
  #pos { color:#6b7280; font-size:13px; }
  main { max-width:1000px; margin:0 auto; padding:22px 24px 40px; }
  .flow { display:none; }
  .flow.on { display:block; }
  .qflow { display:none; }
  .qflow.on { display:block; }
  .tabs { display:flex; border:1px solid #d1d5db; border-radius:8px; overflow:hidden; margin-left:auto; }
  .tab { border:0; background:#fff; padding:7px 16px; font-size:14px; cursor:pointer; color:#374151;
         font-family:inherit; }
  .tab.on { background:#111827; color:#fff; }
  .oflow { display:none; }
  .oflow.on { display:block; }
  .panel-quant, .panel-overlap, .panel-cons, .panel-trace { display:none; }
  body.view-trace .panel-trace { display:block; }
  .fmsec, .trsec { display:none; }
  .fmsec.on, .trsec.on { display:block; }
  .pick-row { display:flex; gap:10px; align-items:center; margin:0 0 12px; font-size:14px; }
  .pick-row select { font-size:14px; padding:6px 10px; min-width:300px; border:1px solid #d1d5db; border-radius:6px; }
  .fm-caption { color:#6b7280; font-size:13px; margin:0 0 10px; }
  body.view-quant .panel-quant { display:block; }
  body.view-overlap .panel-overlap { display:block; }
  body.view-cons .panel-cons { display:block; }
  body.view-quant .panel-flow, body.view-overlap .panel-flow, body.view-cons .panel-flow, body.view-trace .panel-flow { display:none; }
  body.view-cons #pick, body.view-cons .nav, body.view-cons #pos, body.view-trace #pick, body.view-trace .nav, body.view-trace #pos { display:none; }
  .qcards { display:grid; grid-template-columns:repeat(auto-fit,minmax(170px,1fr)); gap:10px; margin:0 0 16px; }
  .qcard { background:#f3f4f6; border-radius:8px; padding:10px 14px; border-top:4px solid #2a9d8f; }
  .qcard .k { font-size:13px; color:#6b7280; }
  .qcard .v { font-size:24px; font-weight:600; margin-top:4px; word-break:break-word; }
  .qgrid { display:grid; grid-template-columns:minmax(220px,1fr) 2fr; gap:16px; margin-bottom:16px; }
  @media (max-width:760px) { .qgrid { grid-template-columns:1fr; } }
  .qbox { border:1px solid #e5e7eb; border-radius:8px; padding:14px 16px; margin-bottom:16px; }
  .qbox h3 { margin:0 0 10px; font-size:16px; }
  .qbox p { margin:4px 0; font-size:14px; }
  .brow { display:grid; grid-template-columns:200px 1fr 70px; gap:10px; align-items:center; margin:7px 0; font-size:14px; }
  .btrack { background:#f3f4f6; border-radius:4px; height:16px; }
  .bbar { background:#2a9d8f; height:16px; border-radius:4px; }
  .bval { text-align:right; color:#374151; font-variant-numeric:tabular-nums; }
  .note { color:#6b7280; font-size:12.5px; margin:8px 0 0; }
  table.qt { width:100%; border-collapse:collapse; font-size:14px; }
  table.qt th, table.qt td { text-align:left; padding:7px 8px; border-bottom:1px solid #e5e7eb; }
  table.qt th { background:#f9fafb; color:#374151; font-weight:600; }
  table.qt td.num { text-align:right; font-variant-numeric:tabular-nums; }
  .warn { background:#fdf0ec; border-left:4px solid #e76f51; padding:10px 14px; border-radius:6px; font-size:14px; }
  h2 { font-size:30px; margin:4px 0 8px; }
  .meta { display:flex; gap:18px; align-items:center; flex-wrap:wrap; margin-bottom:6px; }
  .badge { padding:3px 12px; border-radius:12px; font-weight:600; font-size:14px; }
  .amt { font-size:15px; color:#374151; }
  .amt b { font-size:17px; }
  .src { color:#6b7280; font-size:13px; margin-bottom:12px; }
  .counts { display:flex; gap:8px; flex-wrap:wrap; margin-bottom:18px; }
  .count { background:#f3f4f6; border-radius:8px; padding:6px 12px; font-size:13px; color:#374151;
           border-top:3px solid #9ca3af; }
  .count b { font-size:16px; margin-right:4px; }
  .hint { color:#6b7280; font-size:13px; margin:0 0 10px; }
  @media print { header { position:static; } .nav, #pos { display:none; } }
</style></head><body>
<header>
  <h1>Adjustment Flow Overview</h1>
  <span class="nav" id="prev" title="Previous">&lsaquo;</span>
  <select id="pick">__OPTIONS__</select>
  <span class="nav" id="next" title="Next">&rsaquo;</span>
  <span id="pos"></span>
  <div class="tabs">
    <button class="tab on" data-view="flow">Flow overview</button>
    <button class="tab" data-view="quant">Quantification</button>
    <button class="tab" data-view="overlap">Overlap</button>
    <button class="tab" data-view="cons">Overall flow</button>
    <button class="tab" data-view="trace">Sheet trace</button>
  </div>
</header>
<main><div class="panel-flow">__FLOWS__</div><div class="panel-quant">__QUANT__</div><div class="panel-overlap">__OVERLAP__</div><div class="panel-cons">__CONS__</div><div class="panel-trace">__TRACE__</div></main>
<script>__JS__
  var pick = document.getElementById("pick"), flows = document.querySelectorAll(".flow"),
      qflows = document.querySelectorAll(".qflow"), oflows = document.querySelectorAll(".oflow");
  function setView(v) {
    document.body.classList.remove("view-flow", "view-quant", "view-overlap", "view-cons", "view-trace");
    document.body.classList.add("view-" + v);
    onViewChange(v);
    document.querySelectorAll(".tab").forEach(function (t) {
      t.classList.toggle("on", t.getAttribute("data-view") === v);
    });
  }
  document.querySelectorAll(".tab").forEach(function (t) {
    t.addEventListener("click", function () { setView(t.getAttribute("data-view")); });
  });
  function show(i) {
    i = (i + flows.length) % flows.length;
    flows.forEach(function (f, k) { f.classList.toggle("on", k === i); });
    qflows.forEach(function (f, k) { f.classList.toggle("on", k === i); });
    oflows.forEach(function (f, k) { f.classList.toggle("on", k === i); });
    pick.value = String(i);
    document.getElementById("pos").textContent = (i + 1) + " of " + flows.length;
    try { history.replaceState(null, "", "#adj=" + i); } catch (e) {}
  }
  pick.addEventListener("change", function () { show(+pick.value); });
  document.getElementById("prev").addEventListener("click", function () { show(+pick.value - 1); });
  document.getElementById("next").addEventListener("click", function () { show(+pick.value + 1); });
  var m = /adj=(\\d+)/.exec(location.hash);
  show(m && +m[1] < flows.length ? +m[1] : 0);
</script>
</body></html>"""


def quant_section_html(flow, quant):
    """The Quantification tab for one adjustment, as plain HTML (no plotly, works offline)."""
    if quant is None:
        return '<p class="warn">No quantification file was uploaded, so there is nothing to show here.</p>'
    matched, how = match_adjustment(flow, quant)
    if matched is None:
        names = ", ".join(sorted({r["adjustment"] for r in quant["rows"]}))
        return (f'<p class="warn"><b>{esc(flow["name"])}</b> isn\'t in the quantification file. '
                f'Its adjustments are: {esc(names)}</p>')
    rows = [r for r in quant["rows"] if r["adjKey"] == norm(matched)]
    q = summarise(rows, quant["dims_present"])

    assigned_text = ", ".join(q["assigned"]) or "-"
    assigned_short = assigned_text
    if len(q["assigned"]) > 2:
        assigned_short = f"{q['assigned'][0]}, {q['assigned'][1]} +{len(q['assigned']) - 2}"
    cards = [("Total records", fmt_num(q["records"])), ("Total data points", fmt_num(q["datapoints"])),
             ("Number of sheets", str(q["sheets"])), ("Assigned to", assigned_short)]
    card_html = "".join(f'<div class="qcard"><div class="k">{k}</div><div class="v">{esc(v)}</div></div>'
                        for k, v in cards)
    note = (f' <span class="note">(similar name to {esc(flow["name"])} - please check it is the right one)</span>'
            if how == "similar" else "")

    summary = ('<div class="qbox"><h3>Quantification summary</h3>'
               f'<p><b>Assigned to:</b> {esc(assigned_text)}</p>'
               f'<p><b>Source systems:</b> {esc(", ".join(q["sources"]) or "-")}</p>'
               f'<p><b>Total records:</b> {fmt_num(q["records"])}</p>'
               f'<p><b>Total data points:</b> {fmt_num(q["datapoints"])}</p>'
               f'<p><b>Number of sheets:</b> {q["sheets"]}</p></div>')

    top = max([abs(v) for v in q["dims"].values()] + [0]) or 1
    bar_rows = []
    for d in DIMENSIONS:
        v = q["dims"].get(d) or 0.0
        label = d if d in quant["dims_present"] else f"{d} (not in file)"
        width = abs(v) / top * 100
        bar = f'<div class="bbar" style="width:{width:.1f}%"></div>' if v else ""
        bar_rows.append(f'<div class="brow"><div>{esc(label)}</div><div class="btrack">{bar}</div>'
                        f'<div class="bval">{fmt_num(v) if v else "0"}</div></div>')
    bits = []
    if q["zero_dims"]:
        bits.append("zero: " + ", ".join(q["zero_dims"]))
    if q["missing_dims"]:
        bits.append("not in the file: " + ", ".join(q["missing_dims"]))
    dims_note = f'<p class="note">No bar shown - {esc(" · ".join(bits))}.</p>' if bits else ""
    dims_box = ('<div class="qbox"><h3>Breakdown by dimension</h3>' + "".join(bar_rows)
                + dims_note + "</div>")

    trs = "".join(
        f'<tr><td>{esc(r["File Name"])}</td><td>{esc(r["Sheet Name"])}</td><td>{esc(r["Assigned To"])}</td>'
        f'<td class="num">{fmt_num(r["Total Records"])}</td><td class="num">{fmt_amount(r["Grand Total"])}</td></tr>'
        for _, r in sheet_table(rows).iterrows())
    sheets_box = ('<div class="qbox"><h3>Sheet-level quantification</h3><table class="qt"><thead><tr>'
                  '<th>File Name</th><th>Sheet Name</th><th>Assigned To</th>'
                  '<th class="num">Total Records</th><th class="num">Grand Total</th></tr></thead>'
                  f'<tbody>{trs}</tbody></table></div>')

    return (f'<div class="qcards">{card_html}</div>'
            f'<p class="hint">{len(rows)} quantification row(s) for <b>{esc(matched)}</b>.{note}</p>'
            f'<div class="qgrid">{summary}{dims_box}</div>{sheets_box}')


@st.cache_data(show_spinner=False)
def overview_html(flows, quant=None):
    """Standalone light page: the Flow overview and the Quantification tab for every adjustment."""
    options, sections, qsections, osections = [], [], [], []
    ov_index = overlap_index(flows)
    for i, f in enumerate(flows):
        options.append(f'<option value="{i}">{esc(f["name"])}'
                       + (f' ({esc(f["adjType"])})' if f["adjType"] else "") + "</option>")
        meta = []
        if f["adjType"]:
            bg, fg = TYPE_BADGE.get(f["adjType"], ("#f3f4f6", "#4b5563"))
            meta.append(f'<span class="badge" style="background:{bg};color:{fg}">'
                        f'{esc(f["adjType"])} adjustment</span>')
        if f["amount"] is not None:
            meta.append(f'<span class="amt">{esc(f["amountLabel"])}: <b>{fmt_amount(f["amount"])}</b></span>')
        counts = {b: sum(1 for s in f["steps"] if s["bucket"] == b) for b in af.BUCKET_ORDER}
        shown = BUCKETS + [bk for bk in OPTIONAL_BUCKETS if counts[bk[0]]]
        chips = [f'<span class="count" style="border-top-color:#FBCE07"><b>{len(f["steps"])}</b>Steps</span>'] + \
                [f'<span class="count" style="border-top-color:{c}"><b>{counts[b]}</b>{esc(lbl)}</span>'
                 for b, lbl, c in shown]
        svg, _ = static_flow_svg(f, uid=f"ah{i}")
        sections.append(
            f'<section class="flow"><h2>{esc(f["name"])}</h2>'
            + (f'<div class="meta">{"".join(meta)}</div>' if meta else "")
            + f'<div class="src">{esc(f["file"])} › {esc(f["sheet"])} · from row {f["headerRow"]}</div>'
            + f'<div class="counts">{"".join(chips)}</div>'
            + '<p class="hint">Each arrow goes to the next step. The pill shows the type; '
              '📄 is the work paper. Hover a card for its full details.</p>'
            + f'{svg}</section>')
        qsections.append(f'<section class="qflow"><h2>{esc(f["name"])}</h2>'
                         + quant_section_html(f, quant) + '</section>')
        ov_html, _ = overlap_html(i, flows, ov_index, uid=f"ov{i}")
        osections.append(f'<section class="oflow"><h2>{esc(f["name"])}</h2>' + ov_html + '</section>')
    tr_html, tr_json = trace_sections(flows)
    return (OVERVIEW_PAGE.replace("__OPTIONS__", "".join(options))
            .replace("__FLOWS__", "".join(sections))
            .replace("__QUANT__", "".join(qsections))
            .replace("__OVERLAP__", "".join(osections))
            .replace("__CONS__", fm_sections(flows))
            .replace("__TRACE__", tr_html)
            .replace("__JS__", PAGE_JS.replace("__TRDATA__", tr_json))
            .replace("__VIS__", af.get_vis_js()))


@st.cache_data(show_spinner=False)
def page_html(flow, embedded):
    return af.build_html([flow] if embedded else flow, embedded=embedded)


# ---------------------------------------------------------------- sidebar
with st.sidebar:
    st.header("Adjustment Flow")
    uploads = st.file_uploader("Excel file(s)", type=["xlsx", "xlsm"], accept_multiple_files=True)
    sheet = st.text_input("Only this sheet (optional)", placeholder="exact tab name")
    quant_file = st.file_uploader(
        "Quantification file (optional)", type=["xlsx", "xlsm"], key="quant_upload",
        help="The quantification workbook. Adds metrics, a dimension chart and a sheet "
             "table for the selected adjustment.")

if not uploads:
    st.title("Adjustment Flow Dashboard")
    st.info("Upload one or more Excel files in the sidebar to begin.")
    st.markdown(
        "Each adjustment starts with a details row (**Temp/Perm, name, Return (CY) + amount**), "
        "followed by up to four tables - **Input, Process, Output - Internal, Output - External** - "
        "each with the columns **Sequence, Name, Type, Source, Classification Basis, "
        "Work Paper Name, Work Paper Sheet** (outputs use **Process** and **Format / Basis**).\n\n"
        "Optionally add the **quantification file** in the sidebar to see record counts, data "
        "points, a dimension chart and a sheet table for each adjustment.")
    st.stop()

flows = parse(tuple((u.name, u.getvalue()) for u in uploads), sheet.strip())
if not flows:
    st.error("No adjustments found. Each table needs a header row with at least "
             "Sequence and two other known columns.")
    st.stop()

quant, quant_error = None, None
if quant_file is not None:
    try:
        quant = parse_quant(quant_file.getvalue())
    except Exception as e:                       # a bad file must not break the flow page
        quant_error = str(e)

with st.sidebar:
    labels = [f"{f['name']}" + (f" · {f['adjType'][:4]}" if f["adjType"] else "") +
              f"  ({len(f['steps'])} steps" +
              (f", ⚠ {len(f['issues'])})" if f["issues"] else ")") for f in flows]
    idx = st.selectbox("Flow", range(len(flows)), format_func=lambda i: labels[i])
    st.caption(f"{len(flows)} flow(s) found")
    st.divider()
    base_name = Path(uploads[0].name).stem if len(uploads) == 1 else "adjustment_flows"
    st.download_button(
        "⬇ Download Flow overview (HTML)", data=overview_html(flows, quant),
        file_name=base_name + "_overview.html", mime="text/html", use_container_width=True,
        help="The light flow overview for every adjustment, with a dropdown and a "
             "Quantification tab (if the quantification file is uploaded). "
             "Standalone - works offline in Chrome or Edge.")
    st.download_button(
        "⬇ Download Interactive view (HTML)", data=page_html(flows, False),
        file_name=base_name + "_flow.html",
        mime="text/html", use_container_width=True,
        help="The dark interactive view for every adjustment. Standalone - works offline in Chrome or Edge.")

flow = flows[idx]
steps = flow["steps"]

# ---------------------------------------------------------------- header + cards
st.title(flow["name"])
badges = []
if flow["adjType"]:
    bg, fg = TYPE_BADGE.get(flow["adjType"], ("#f3f4f6", "#4b5563"))
    badges.append(f'<span style="background:{bg};color:{fg};padding:3px 12px;border-radius:12px;'
                  f'font-weight:600;font-size:0.9rem">{esc(flow["adjType"])} adjustment</span>')
if flow["amount"] is not None:
    badges.append(f'<span style="font-size:0.95rem;color:#374151">{esc(flow["amountLabel"])}: '
                  f'<b style="font-size:1.1rem">{fmt_amount(flow["amount"])}</b></span>')
if badges:
    st.markdown('<div style="display:flex;gap:18px;align-items:center;margin:-6px 0 6px">'
                + "".join(badges) + "</div>", unsafe_allow_html=True)
st.caption(f"{flow['file']} › {flow['sheet']} · from row {flow['headerRow']}")

counts = {b: sum(1 for s in steps if s["bucket"] == b) for b in af.BUCKET_ORDER}
shown = BUCKETS + [bk for bk in OPTIONAL_BUCKETS if counts[bk[0]]]
cards = [("Steps", len(steps), "#FBCE07")] + \
        [(label, counts[b], col) for b, label, col in shown] + \
        [("File links", len(flow["dataEdges"]), "#FBCE07"),
         ("Issues", len(flow["issues"]), "#e74c3c" if flow["issues"] else "#2ecc71")]
for col, (label, value, _) in zip(st.columns(len(cards)), cards):
    col.metric(label, value)
# colour each card's top border
st.markdown("<style>" + "".join(
    f'div[data-testid="stHorizontalBlock"] > div:nth-child({i + 1}) div[data-testid="stMetric"]'
    f'{{border-top-color:{c};}}' for i, (_, _, c) in enumerate(cards)) + "</style>",
    unsafe_allow_html=True)

# ---------------------------------------------------------------- 1. static flow
st.subheader("Flow overview")
st.caption("Steps in order, grouped by stage - each arrow goes to the next step. "
           "The pill shows the type; 📄 is the file used. Hover a card for its full details.")
svg, svg_h = static_flow_svg(flow)
components.html(f'<div style="display:flex;justify-content:center">{svg}</div>',
                height=int(svg_h) + 12, scrolling=False)

# ---------------------------------------------------------------- 2. interactive flow
st.divider()
st.subheader("Interactive view")
st.caption("Click a step for its details on the right. Tick **Step-by-step only** "
           "to see just the 1 → 2 → 3 order.")
components.html(page_html(flow, True), height=780, scrolling=False)

# ---------------------------------------------------------------- 3. quantification
st.divider()
st.subheader("Quantification")

if quant_file is None:
    st.info("Upload the quantification file in the sidebar to see this adjustment's record "
            "count, data points, dimension breakdown and sheet list.")
elif quant_error:
    st.error(f"Couldn't read the quantification file: {quant_error}")
elif not quant["found_header"]:
    st.warning("No quantification header row found. The sheet needs a row with **Adjustment** "
               "and at least three of: File Name, Sheet Name, Assigned To, Source System, "
               "Total Records, the dimension columns or Grand Total.")
elif not quant["rows"]:
    st.warning("The quantification header was found, but it has no adjustment rows under it.")
else:
    matched, how = match_adjustment(flow, quant)
    if matched is None:
        names = sorted({r["adjustment"] for r in quant["rows"]})
        st.warning(f"**{flow['name']}** isn't in the quantification file. Its adjustments are: "
                   + ", ".join(names))
    else:
        rows = [r for r in quant["rows"] if r["adjKey"] == norm(matched)]
        q = summarise(rows, quant["dims_present"])
        note = f" (similar name to **{flow['name']}** - please check it is the right one)" \
            if how == "similar" else ""
        st.caption(f"{len(rows)} quantification row(s) for **{matched}**{note}")

        # metric cards, styled like the cards above
        assigned_text = ", ".join(q["assigned"]) if q["assigned"] else "-"
        if len(q["assigned"]) > 2:
            assigned_text = f"{q['assigned'][0]}, {q['assigned'][1]} +{len(q['assigned']) - 2}"
        with st.container(key="quant_cards"):
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Total records", fmt_num(q["records"]))
            m2.metric("Total data points", fmt_num(q["datapoints"]),
                      help="Sum of the Grand Total column for this adjustment.")
            m3.metric("Number of sheets", q["sheets"])
            m4.metric("Assigned to", assigned_text,
                      help=", ".join(q["assigned"]) if q["assigned"] else "No name in the file.")

        left, right = st.columns([1, 2], gap="large")
        with left:
            with st.container(border=True):
                st.markdown("**Quantification summary**")
                st.markdown(
                    f"**Assigned to:** {', '.join(q['assigned']) or '-'}  \n"
                    f"**Source systems:** {', '.join(q['sources']) or '-'}  \n"
                    f"**Total records:** {fmt_num(q['records'])}  \n"
                    f"**Total data points:** {fmt_num(q['datapoints'])}  \n"
                    f"**Number of sheets:** {q['sheets']}")
        with right:
            with st.container(border=True):
                st.markdown("**Breakdown by dimension**")
                if go is None:
                    st.error("The chart needs plotly. Run: `pip install plotly`")
                else:
                    st.plotly_chart(dimension_figure(q["dims"], quant["dims_present"]),
                                    use_container_width=True, config={"displayModeBar": False})
                bits = []
                if q["zero_dims"]:
                    bits.append("zero: " + ", ".join(q["zero_dims"]))
                if q["missing_dims"]:
                    bits.append("not in the file: " + ", ".join(q["missing_dims"]))
                if bits:
                    st.caption("No bar shown - " + " · ".join(bits) + ".")

        st.markdown("**Sheet-level quantification**")
        st.dataframe(
            sheet_table(rows), use_container_width=True, hide_index=True,
            column_config={
                "Total Records": st.column_config.NumberColumn(format="%.0f"),
                "Grand Total": st.column_config.NumberColumn(format="%.2f"),
            })

# ---------------------------------------------------------------- 4. overlap
st.divider()
st.subheader("Overlap with other adjustments")
st.caption("A step overlaps when another adjustment uses the same work paper (file and sheet). "
           "Shared steps have a dashed purple outline; the table shows where else each one is used.")
ov_index = overlap_index(flows)
ov_html, ov_h = overlap_html(idx, flows, ov_index, uid="ovs")
components.html(ov_html, height=ov_h, scrolling=True)

# ---------------------------------------------------------------- 5. overall flow
st.divider()
st.subheader("Overall flow")
st.caption("Inputs on top, one process lane per adjustment, then the internal and external outputs. "
           "Click a card or lane to lock its connections. Pick one adjustment to see only its flow.")
fm_names = ["All adjustments"] + [f["name"] for f in flows]
fm_pick = st.selectbox("Show", range(len(fm_names)), format_func=lambda i: fm_names[i], key="fm_pick")
map_html, map_h = overall_flow_html(flows, None if fm_pick == 0 else fm_pick - 1, uid="fms")
components.html(map_html + f"<script>{FM_JS}</script>", height=map_h, scrolling=True)

st.divider()
st.subheader("Sheet trace view")
st.caption("The same dataflow in the sheet-trace layout: inputs at the top, adjustments in the middle, "
           "outputs at the bottom. Search, filter by arrow counts, click a box to highlight its links.")
tr_pick = st.selectbox("Show ", range(len(fm_names)), format_func=lambda i: fm_names[i], key="trace_pick")
components.html(trace_doc(flows, None if tr_pick == 0 else tr_pick - 1), height=760, scrolling=False)

# ---------------------------------------------------------------- 6. tables
st.divider()
tab_steps, tab_issues = st.tabs(["Steps table", f"Issues ({len(flow['issues'])})"])

with tab_steps:
    rows = []
    for s in steps:
        # one set of columns for every stage (outputs call Name "Process"
        # and Classification Basis "Format / Basis")
        row = {"Stage": s["stage"], "Sequence": f"Step {s['step']}" if s["step"] else "",
               "Name / Process": s["report"], "Type": s["action"], "Source": s["source"],
               "Classification / Format Basis": s["standard"],
               "Work Paper Name": s["location"], "Work Paper Sheet": s["wpSheet"]}
        row.update(dict(s["extras"]))
        row["Excel row"] = s["row"]
        rows.append(row)
    st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

with tab_issues:
    if not flow["issues"]:
        st.success("No issues found.")
    else:
        by_id = {s["id"]: s for s in steps}
        st.dataframe(pd.DataFrame([{
            "Step": (f"Step {by_id[i['id']]['step'] or '?'} · {by_id[i['id']]['report']}"
                     if i["id"] else "General"),
            "Excel row": by_id[i["id"]]["row"] if i["id"] else "",
            "Issue": i["msg"],
        } for i in flow["issues"]]), use_container_width=True, hide_index=True)
